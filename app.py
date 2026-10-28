import io
import re
from pathlib import Path
import cv2
import joblib
import numpy as np
import streamlit as st
from PIL import Image
from pypdf import PdfReader
from docx import Document

st.set_page_config(page_title="ForensiQ | Digital Forensics", page_icon="🔎", layout="wide")
BASE_DIR = Path(__file__).resolve().parent
MODEL_PATH = BASE_DIR / "models" / "forensiq_model.joblib"
IMAGE_SIZE = (64, 64)

st.markdown("""
<style>
.stApp {background: radial-gradient(ellipse at 15% 0%, #142844 0%, transparent 42%), #080D17; color:#EAF1FC;}
[data-testid="stSidebar"] {background:#0C1422;border-right:1px solid #23344B;}
.block-container {max-width:1200px;padding-top:2rem;}
.hero {padding:26px;border:1px solid #263B58;border-radius:18px;background:linear-gradient(120deg,#13243A,#101827 70%);margin-bottom:22px;}
.eyebrow {color:#62D8FF;font-size:11px;font-weight:800;letter-spacing:3px;}
.hero h1 {color:#F5F8FF;font-size:38px;margin:10px 0 8px 0;}
.hero p,.muted {color:#A7B8CE;line-height:1.7;}
.panel,.result {background:#101A2A;border:1px solid #273951;border-radius:14px;padding:18px;margin-bottom:12px;}
.result {background:#10283A;border-color:#2A526D;margin-top:16px;}
div.stButton > button {background:linear-gradient(100deg,#2478F5,#20BDEB);color:white;border:0;border-radius:10px;min-height:44px;font-weight:700;}
div[data-testid="stFileUploader"] {background:#101B2C;border-radius:12px;border:1px solid #283D58;padding:10px;}
</style>
""", unsafe_allow_html=True)

@st.cache_resource
def load_model_bundle():
    if not MODEL_PATH.exists():
        return None
    try:
        return joblib.load(MODEL_PATH)
    except Exception:
        return None

def prepare_image(image: Image.Image) -> np.ndarray:
    arr = np.asarray(image.convert("RGB"))
    resized = cv2.resize(arr, IMAGE_SIZE)
    return resized.astype(np.float32).flatten().reshape(1, -1) / 255.0

def predict_image(image: Image.Image, bundle):
    model = bundle["model"]
    features = prepare_image(image)
    prediction = int(model.predict(features)[0])
    scores = None
    if hasattr(model, "predict_proba"):
        try:
            scores = model.predict_proba(features)[0]
        except Exception:
            pass
    label = "Forged" if prediction == 1 else "Authentic"
    score = float(scores[prediction]) * 100 if scores is not None and len(scores) > prediction else None
    return label, score

def document_amount_check(text: str):
    """
    Rule-based arithmetic consistency check for readable invoices.
    This can flag obvious amount mismatches, but is not a general forgery detector.
    """
    normalized = text.replace(",", "")
    lines = [re.sub(r"\s+", " ", line).strip() for line in normalized.splitlines() if line.strip()]

    def numbers(line):
        # Currency glyphs may be lost during PDF extraction, so extract numeric
        # values independently of the currency symbol.
        return [float(x) for x in re.findall(r"(?<![A-Za-z])\d+(?:\.\d{1,2})?(?![A-Za-z])", line)]

    # Check table rows such as "Office notebooks | 10 | ₹350 | ₹8,500".
    for line in lines:
        cells = [c.strip() for c in line.split("|")]
        if len(cells) >= 4:
            q, u, total = numbers(cells[1]), numbers(cells[2]), numbers(cells[3])
            if q and u and total and q[-1] > 0 and u[-1] > 0 and total[-1] > 0:
                if abs(q[-1] * u[-1] - total[-1]) > 0.01:
                    return ("Amount consistency",
                            f"Possible line-item mismatch: quantity {q[-1]:g} × unit price {u[-1]:g} = {q[-1]*u[-1]:g}, but listed total is {total[-1]:g}.",
                            "Review")

    # Check common PDF extraction where table cells appear on separate lines:
    # Description / Quantity / Unit price / Total / item / 10 / 350 / 8500
    for i, line in enumerate(lines):
        if re.fullmatch(r"quantity|qty", line, re.I):
            nearby = lines[i + 1:i + 9]
            labels = [x.lower() for x in nearby]
            try:
                unit_idx = next(j for j, x in enumerate(labels) if "unit price" in x or x == "rate")
                total_idx = next(j for j, x in enumerate(labels) if x == "total" or "line total" in x)
                # Quantity and unit price labels should occur before the data rows.
                data_start = max(unit_idx, total_idx) + 1
                values = []
                for candidate in nearby[data_start:]:
                    vals = numbers(candidate)
                    if vals:
                        values.append(vals[-1])
                    if len(values) == 3:
                        break
                if len(values) >= 3:
                    qty, unit, listed_total = values[:3]
                    if qty > 0 and unit > 0 and listed_total > 0 and abs(qty * unit - listed_total) > 0.01:
                        return ("Amount consistency",
                                f"Possible line-item mismatch: quantity {qty:g} × unit price {unit:g} = {qty*unit:g}, but listed total is {listed_total:g}.",
                                "Review")
            except StopIteration:
                pass

    # Compare multiple explicit total labels when they disagree.
    total_values = []
    for line in lines:
        if re.search(r"\b(grand\s*total|amount\s*due|total\s*payable|invoice\s*total|total\s*amount)\b", line, re.I):
            vals = numbers(line)
            if vals:
                total_values.append(vals[-1])
    if len(set(round(v, 2) for v in total_values)) > 1:
        return ("Amount consistency", f"Different values appear beside total labels: {total_values}", "Review")

    return ("Amount consistency", "No obvious mismatch found by the basic amount checks", "Info")

def analyze_pdf(file_bytes: bytes):
    reader = PdfReader(io.BytesIO(file_bytes))
    metadata = {}
    if reader.metadata:
        for key in ("/Title", "/Author", "/Creator", "/Producer", "/CreationDate", "/ModDate", "/Subject"):
            value = reader.metadata.get(key)
            if value:
                metadata[key.lstrip("/")] = str(value)[:300]
    page_text_lengths, parts = [], []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        page_text_lengths.append(len(text.strip()))
        if text.strip():
            parts.append(text.strip())
    text = "\n\n".join(parts)
    checks = [
        ("Metadata", "Some document metadata is present" if metadata else "No common metadata fields found", "Info" if metadata else "Review"),
        ("Text extraction", f"Extracted {len(text)} text characters" if text else "No readable text extracted; it may be scanned or image-only", "Info" if text else "Review"),
        ("Encryption", "PDF reports encryption" if reader.is_encrypted else "PDF is not marked as encrypted", "Review" if reader.is_encrypted else "Info"),
        document_amount_check(text),
    ]
    return {"pages": len(reader.pages), "metadata": metadata, "text": text, "page_text_lengths": page_text_lengths, "checks": checks}

def analyze_docx(file_bytes: bytes):
    document = Document(io.BytesIO(file_bytes))
    props = document.core_properties
    metadata = {}
    for label, value in [("Title", props.title), ("Author", props.author), ("Subject", props.subject),
                         ("Created", props.created), ("Modified", props.modified),
                         ("Last modified by", props.last_modified_by), ("Revision", props.revision)]:
        if value not in (None, ""):
            metadata[label] = str(value)[:300]
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    table_rows = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                table_rows.append(" | ".join(cells))
    combined = "\n".join(paragraphs + table_rows)
    checks = [
        ("Metadata", "Document properties are present" if metadata else "Few or no document properties found", "Info" if metadata else "Review"),
        ("Text content", f"Extracted {len(combined)} characters from paragraphs and tables" if combined else "No text found in paragraphs or tables", "Info" if combined else "Review"),
        ("Tables", f"Found {len(document.tables)} table(s)", "Info"),
        document_amount_check(combined),
    ]
    return {"metadata": metadata, "text": combined, "paragraph_count": len(paragraphs), "table_count": len(document.tables), "checks": checks}

if "page_choice" not in st.session_state:
    st.session_state["page_choice"] = "Dashboard"

def go_to(page_name):
    st.session_state["page_choice"] = page_name

with st.sidebar:
    st.markdown("<div style='font-size:25px;font-weight:800;'>🔎 Forensi<span style='color:#58D5FF;'>Q</span></div>", unsafe_allow_html=True)
    st.caption("DIGITAL FORENSICS")
    page = st.radio("MENU", ["Dashboard", "Image Forensics", "Document Forensics"], key="page_choice")
    st.markdown("---")
    st.caption("Prototype • Preliminary screening only")

if page == "Dashboard":
    st.markdown("""<div class="hero"><div class="eyebrow">DIGITAL AUTHENTICITY LAB</div><h1>ForensiQ</h1><p>Inspect images and documents in one simple workspace.</p></div>""", unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        st.markdown('<div class="panel"><h3>◈ Image Forensics</h3><p class="muted">Upload an image and view the experimental model prediction.</p></div>', unsafe_allow_html=True)
        st.button("Open Image Forensics", use_container_width=True, on_click=go_to, args=("Image Forensics",))
    with c2:
        st.markdown('<div class="panel"><h3>▤ Document Forensics</h3><p class="muted">Inspect PDF and Word document properties and content.</p></div>', unsafe_allow_html=True)
        st.button("Open Document Forensics", use_container_width=True, on_click=go_to, args=("Document Forensics",))
    st.info("Prototype only. Automated checks are indicators for review, not proof of authenticity or forgery.")

elif page == "Image Forensics":
    st.markdown("""<div class="hero"><div class="eyebrow">IMAGE ANALYSIS</div><h1>Image Forensics</h1><p>Upload an image and view the model's experimental prediction.</p></div>""", unsafe_allow_html=True)
    bundle = load_model_bundle()
    if bundle is None:
        st.error(f"Trained model not found or could not be loaded. Expected file: {MODEL_PATH}")
    else:
        uploaded = st.file_uploader("Choose an image", type=["jpg", "jpeg", "png", "bmp", "tif", "tiff"], key="image_upload")
        if uploaded:
            try:
                image = Image.open(io.BytesIO(uploaded.getvalue())).convert("RGB")
                left, right = st.columns([1, 1], gap="large")
                with left:
                    st.image(image, caption=uploaded.name, use_container_width=True)
                    st.caption(f"Dimensions: {image.width} × {image.height} pixels")
                with right:
                    st.markdown("### Result")
                    if st.button("Analyze Image", type="primary", use_container_width=True):
                        with st.spinner("Analyzing..."):
                            label, score = predict_image(image, bundle)
                        st.session_state["last_image_result"] = {"name": uploaded.name, "label": label, "score": score}
                result = st.session_state.get("last_image_result")
                if result and result.get("name") == uploaded.name:
                    st.markdown('<div class="result">', unsafe_allow_html=True)
                    st.subheader(f"Model prediction: {result['label']}")
                    if result["score"] is not None:
                        st.write(f"Model score: {result['score']:.1f}% (not a verified probability)")
                    st.markdown("</div>", unsafe_allow_html=True)
                    st.caption("Experimental result only; this baseline may be wrong and is not forensic proof.")
            except Exception as exc:
                st.error(f"Could not process this image: {exc}")

elif page == "Document Forensics":
    st.markdown("""<div class="hero"><div class="eyebrow">DOCUMENT ANALYSIS</div><h1>Document Forensics</h1><p>Upload a PDF or DOCX file to inspect its properties and extract readable text.</p></div>""", unsafe_allow_html=True)
    document_file = st.file_uploader("Choose a document", type=["pdf", "docx"], key="document_upload_v2",
                                     help="PDF and DOCX supported. Legacy .doc files are not supported directly.")
    if document_file:
        file_bytes = document_file.getvalue()
        st.caption(f"File: {document_file.name} • Size: {len(file_bytes) / 1024:.1f} KB")
        if st.button("Analyze Document", type="primary", use_container_width=True):
            try:
                with st.spinner("Inspecting document..."):
                    suffix = Path(document_file.name).suffix.lower()
                    result = analyze_pdf(file_bytes) if suffix == ".pdf" else analyze_docx(file_bytes)
                    result_type = suffix[1:].upper()
                st.session_state["last_document_result"] = {"name": document_file.name, "type": result_type, "result": result}
            except Exception as exc:
                st.error(f"Could not analyze this document: {exc}")
        saved = st.session_state.get("last_document_result")
        if saved and saved.get("name") == document_file.name:
            result = saved["result"]
            st.markdown("### Document summary")
            if saved["type"] == "PDF":
                c1, c2 = st.columns(2)
                c1.metric("Pages", result["pages"])
                c2.metric("Pages with extractable text", sum(1 for n in result["page_text_lengths"] if n > 0))
            else:
                c1, c2, c3 = st.columns(3)
                c1.metric("Paragraphs", result["paragraph_count"])
                c2.metric("Tables", result["table_count"])
                c3.metric("Extracted characters", len(result["text"]))
            amount_flags = [item for item in result["checks"] if item[0] == "Amount consistency" and item[2] == "Review"]
            review_flags = [item for item in result["checks"] if item[2] == "Review"]
            st.markdown("### Analysis Result")
            if amount_flags:
                st.error("POTENTIALLY SUSPICIOUS — an amount inconsistency was detected.")
            elif review_flags:
                st.warning("REVIEW REQUIRED — basic checks found something that needs inspection.")
            else:
                st.success("NO OBVIOUS WARNING SIGNS — this does not prove the document is authentic.")
            st.markdown("### Basic checks")
            for name, detail, status in result["checks"]:
                if status == "Review":
                    st.warning(f"**{name}:** {detail}")
                else:
                    st.success(f"**{name}:** {detail}")
            if result["metadata"]:
                with st.expander("View document metadata"):
                    for key, value in result["metadata"].items():
                        st.write(f"**{key}:** {value}")
            else:
                st.caption("No common metadata fields were found.")
            if result["text"].strip():
                with st.expander("View extracted text"):
                    st.text_area("Extracted content", result["text"][:20000], height=260, label_visibility="collapsed")
            st.warning("Screening only: missing metadata, encryption, or little extractable text does not prove forgery. This version does not classify documents as authentic or forged.")
    st.caption("Supported: PDF and DOCX. Older .doc files must be converted to .docx first.")

st.markdown("---")
st.markdown("<div style='text-align:center;color:#7286A0;font-size:11px;padding:10px;'>FORENSIQ • IMAGE & DOCUMENT FORENSICS<br><br>Prototype only • Results require validation and human review</div>", unsafe_allow_html=True)
